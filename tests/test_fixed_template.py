import copy
import importlib.util
import tempfile
import unittest
import json
import hashlib
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'skills/analog-recovery-report/scripts/fill_template.py'


class FixedTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location('fixed_template', SCRIPT)
        cls.m = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.m)

    def payload(self):
        return self.m.example_payload()

    def test_missing_section_rejected(self):
        p = self.payload(); del p['sections']['1.5']
        self.assertTrue(self.m.validate(p))

    def test_unknown_reference_rejected(self):
        p = self.payload(); p['sections']['1.5']['paragraphs'][0]['refs'] = ['missing']
        self.assertTrue(self.m.validate(p))

    def test_unmapped_port_rejected(self):
        p = self.payload(); p['expected_ports'].append('NEW_OUTPUT')
        self.assertTrue(self.m.validate(p))

    def test_manual_numbering_rejected(self):
        p = self.payload(); p['sections']['1.1']['paragraphs'][0]['text'] = '1) 重复编号'
        self.assertTrue(self.m.validate(p))

    def test_typo_section_rejected(self):
        p = self.payload(); p['sections']['1.99'] = {}
        self.assertTrue(self.m.validate(p))

    def test_valid_synthetic_package(self):
        self.assertEqual([], self.m.validate(self.payload()))

    def test_condition_rows_are_required_and_not_metrics(self):
        p = self.payload()
        self.assertEqual({'supply_voltage', 'temperature'}, {r['quantity'] for r in p['sections']['1.3']['condition_rows']})
        self.assertTrue(all(len(r['cells']) == 8 and 'metric_ref' not in r for r in p['sections']['1.3']['condition_rows']))
        del p['sections']['1.3']['condition_rows']
        self.assertTrue(any('missing condition_rows' in e for e in self.m.validate(p)))

    def test_supported_architecture_requires_diagram(self):
        p=self.payload();p['sections']['1.5']['status']='supported'
        self.assertTrue(self.m.validate(p))

    def test_simulation_metrics_require_electrical_packet(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td);(p/'metrics.json').write_text(json.dumps({'metrics':[{'samples':[{'run_ref':'run'}]}]}),encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'measurement-packet'):self.m.check_electrical(p,None)

    def test_static_report_does_not_require_simulator(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td);(p/'metrics.json').write_text(json.dumps({'metrics':[{'samples':[]}]}),encoding='utf-8')
            self.assertIsNone(self.m.check_electrical(p,None))

    def test_production_electrical_handoff_reads_actual_artifacts(self):
        from test_measurements_gate import packet
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);p=packet()
            raw=root/'raw.psf';raw.write_text('\n'.join(s['psf_sample'] for s in p['raw_samples']),encoding='utf-8')
            log=root/'run.log';log.write_text(
                'Spectre (R) Circuit Simulator\n'
                'Version 23.1.0\n'
                'Command line:\n'
                'spectre -64 run.scs\n'
                "Simulating `run.scs'\n"
                'spectre completes with 0 errors, 0 warnings, 0 notices.\n',
                encoding='utf-8',
            )
            for s in p['raw_samples']:
                s['source'].update(path=str(raw),sha256=hashlib.sha256(raw.read_bytes()).hexdigest())
            for run in p['runs']:
                run.pop('source',None);run.update(log_path=str(log),sha256=hashlib.sha256(log.read_bytes()).hexdigest())
            path=root/'packet.json';path.write_text(json.dumps(p),encoding='utf-8')
            m=p['measurements'][0]
            (root/'metrics.json').write_text(json.dumps({'metrics':[{'id':m['id'],'unit':m['unit'],'samples':[{'run_ref':p['runs'][0]['id']}]}]}),encoding='utf-8')
            self.assertTrue(self.m.check_electrical(root,path)['valid'])

    def test_wrong_table_cell_type_rejected(self):
        p=self.payload();p['sections']['1.3']['rows'][0]['cells'][0]={'unexpected':'object'}
        self.assertTrue(self.m.validate(p))

    def test_bundle_port_omission_cannot_be_hidden(self):
        from test_skill_handoff import make_bundle
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);bundle=make_bundle(root)
            for name,data in bundle.items():(root/(name+'.json')).write_text(json.dumps(data),encoding='utf-8')
            p=self.payload();p['evidence_ids']=['c_direct'];p['sections']['1.3']['rows']=[]
            for s in p['sections'].values():
                for item in s.get('paragraphs',[])+s.get('rows',[]):item['refs']=['c_direct']
            p['verification'][0]['refs']=['c_direct']
            bound,_=self.m.bind_bundle(p,root,root)
            self.assertTrue(self.m.validate(bound))

    def test_bundle_rejects_invented_evidence_id(self):
        from test_skill_handoff import make_bundle
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);bundle=make_bundle(root)
            for name,data in bundle.items():(root/(name+'.json')).write_text(json.dumps(data),encoding='utf-8')
            with self.assertRaisesRegex(ValueError,'invents evidence'):self.m.bind_bundle(self.payload(),root,root)

    def test_metrics_are_filled_from_samples_not_content_numbers(self):
        from test_skill_handoff import make_bundle
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);bundle=make_bundle(root)
            for name,data in bundle.items():(root/(name+'.json')).write_text(json.dumps(data),encoding='utf-8')
            p=self.payload();p['evidence_ids']=['m_sim']
            row=p['sections']['1.3']['rows'][0];row['metric_ref']='m_sim';row['refs']=['m_sim'];row['cells'][4:8]=['999','999','999','W']
            bound,_=self.m.bind_bundle(p,root,root)
            self.assertEqual(['-1e+308','2.5','1e+308','V'],bound['sections']['1.3']['rows'][0]['cells'][4:8])
            self.assertNotIn('未执行仿真', bound['sections']['1.3']['rows'][0]['cells'][3])
            self.assertIn('min [case_low] low', bound['sections']['1.3']['rows'][0]['cells'][3])
            self.assertIn('max [case_high] high', bound['sections']['1.3']['rows'][0]['cells'][3])
            self.assertTrue(bound['sections']['1.3']['rows'][0]['condition_cases'])

    def test_figure_embedding_and_trace(self):
        from PIL import Image
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            path=Path(td);Image.new('RGB',(50,20),'white').save(path/'figure.png')
            p=self.payload();p['sections']['1.5']['figures']=[{'path':str(path/'figure.png'),'kind':'concept','caption':'独立图片夹具','refs':['synthetic']}]
            self.m.build(p,path/'out.docx')
            self.assertEqual(1,len(Document(path/'out.docx').inline_shapes))
            trace=json.loads((path/'out.trace.json').read_text(encoding='utf-8'))
            self.assertTrue(any(x.get('caption')=='独立图片夹具' for x in trace['blocks']))

    def test_template_change_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td)/'bad.docx'; p.write_bytes(b'changed')
            with self.assertRaises(ValueError): self.m.build(self.payload(), Path(td)/'out.docx', template=p)

    def test_docx_structure_and_clean_body(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as td:
            out=Path(td)/'out.docx'; self.m.build(self.payload(),out)
            d=Document(out)
            self.assertFalse(self.m.audit(out)['errors'])
            texts=[p.text for p in d.paragraphs]
            self.assertIn('实现方案', texts)
            self.assertNotIn('模拟模块2',texts)
            self.assertFalse(any('此处描述' in x for x in texts))
            labels=[p for p in d.paragraphs if p.text.startswith('仿真工具：')]
            self.assertTrue(labels)
            self.assertIsNone(labels[0]._p.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr'))
            with self.assertRaises(FileExistsError): self.m.build(self.payload(),out)

    def test_drawing_property_ids_are_unique_across_document_parts(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as td:
            path=Path(td); Image.new('RGB',(20,20),'white').save(path/'figure.png')
            p=self.payload(); p['sections']['1.5']['figures']=[{'path':str(path/'figure.png'),'kind':'concept','caption':'ID test','refs':['synthetic']}]
            out=path/'out.docx'; self.m.build(p,out)
            ns={'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
            ids=[]
            with zipfile.ZipFile(out) as z:
                for name in z.namelist():
                    if name.endswith('.xml'):
                        root=ET.fromstring(z.read(name))
                        ids.extend(n.get('id') for n in root.findall('.//wp:docPr',ns))
            self.assertEqual(len(ids),len(set(ids)))


if __name__ == '__main__': unittest.main()
