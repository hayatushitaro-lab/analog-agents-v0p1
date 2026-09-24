"""Independent report-contract regressions; no project-specific expected numbers."""
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('filler_v010', ROOT/'skills/analog-recovery-report/scripts/fill_template.py')
f = importlib.util.module_from_spec(spec)
spec.loader.exec_module(f)


class PresentationContractTests(unittest.TestCase):
    def test_plan_is_generated_from_verification(self):
        p = f.example_payload()
        rows = f.verification_rows(p)
        self.assertEqual('N', rows[0]['cells'][2])
        self.assertEqual(p['verification'][0]['title'], rows[0]['cells'][1])

    def test_global_verification_figures_rejected(self):
        p = f.example_payload()
        p['sections']['1.8']['figures'] = [{'path': __file__, 'caption': 'x', 'kind': 'concept', 'refs': ['synthetic']}]
        self.assertTrue(any('verification item' in e for e in f.validate(p)))

    def test_metric_needs_matching_test(self):
        p = f.example_payload()
        p['sections']['1.3']['rows'][0]['metric_ref'] = 'uncovered'
        self.assertTrue(any('uncovered metric' in e for e in f.validate(p)))

    def test_no_ids_in_plan_status(self):
        p = f.example_payload(); p['verification'][0]['status'] = 'm_metric'
        self.assertTrue(f.validate(p))

    def test_duplicate_models_rejected(self):
        p = f.example_payload()
        p['sections']['1.7.2']['rows'] = [{'cells': ['generic_nmos', 'NMOS'], 'refs': ['synthetic']}] * 2
        self.assertTrue(any('duplicate model' in e for e in f.validate(p)))

    def test_declared_plan_cannot_disagree(self):
        p = f.example_payload()
        p['sections']['1.6']['rows'] = [{'cells': ['x', 'unrelated', 'Y', 'x'], 'refs': ['synthetic']}]
        self.assertTrue(any('generated verification plan' in e for e in f.validate(p)))

    def test_figure_placed_in_its_subsection(self):
        from PIL import Image
        from docx import Document
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as td:
            td = Path(td); im = td/'figure.png'; Image.new('RGB', (80, 40), 'white').save(im)
            p = f.example_payload()
            p['verification'][0]['figures'] = [{'path': str(im), 'caption': 'UNIQUE_CAPTION', 'kind': 'concept', 'refs': ['synthetic']}]
            out = td/'out.docx'; f.build(p, out); d = Document(out)
            texts = [x.text for x in d.paragraphs]
            self.assertLess(texts.index(p['verification'][0]['title']), next(i for i,x in enumerate(texts) if 'UNIQUE_CAPTION' in x))
            self.assertTrue(any(t.cell(0, 2).text == '状态' for t in d.tables if len(t.columns) == 4))
            labels = [x for x in d.paragraphs if x.text.startswith('仿真方法：')]
            self.assertIsNone(labels[0]._p.find('.//'+qn('w:numPr')))

    def test_interface_grouping(self):
        rows = [{'cells': [name, direction, '', '', '']} for name, direction in [('Z', '输出'), ('B', '输入'), ('A', '输入')]]
        self.assertEqual(['A', 'B', 'Z'], [r['cells'][0] for r in f.group_interfaces(rows)])

    def test_single_point_remains_visible_without_typ(self):
        from test_skill_handoff import make_bundle
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); bundle = make_bundle(root)
            metric = bundle['metrics']['metrics'][0]
            metric['samples'] = metric['samples'][:1]
            metric['expected_cases'] = metric['expected_cases'][:1]
            metric['typical_case_id'] = None
            for name, value in bundle.items():
                (root/(name+'.json')).write_text(json.dumps(value), encoding='utf-8')
            p = f.example_payload(); p['evidence_ids'] = [metric['id']]
            row = p['sections']['1.3']['rows'][0]
            row.update(metric_ref=metric['id'], refs=[metric['id']])
            bound, _ = f.bind_bundle(p, root, root)
            cells = bound['sections']['1.3']['rows'][0]['cells']
            self.assertIn('已测 ', cells[3])
            self.assertEqual('未指定典型工况', cells[5])
            self.assertEqual(['—', '—'], [cells[4], cells[6]])

    def test_analytical_is_valid_figure_kind(self):
        p = f.example_payload()
        p['verification'][0]['figures'] = [{'path': __file__, 'caption': '解析示意', 'kind': 'analytical', 'refs': ['synthetic']}]
        self.assertEqual([], f.validate(p))
