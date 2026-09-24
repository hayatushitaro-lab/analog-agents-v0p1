"""Role, evidence and document boundary tests using only synthetic artifacts."""
import copy
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from test_skill_handoff import make_bundle, CHECKER
from test_recovery_quality import load

f=load('fill_v09','skills/analog-recovery-report/scripts/fill_template.py')
d=load('docx_v09','skills/analog-recovery-report/scripts/check_docx.py')
r=load('render_v09','skills/analog-recovery-report/scripts/render_report.py')
t=load('tasks_v09','skills/analog-recovery/scripts/task_card.py')

class IntegrationTests(unittest.TestCase):
    def test_partial_keeps_declared_typical(self):
        with tempfile.TemporaryDirectory() as tmp:
            b=make_bundle(Path(tmp));m=b['metrics']['metrics'][0]
            m['samples']=m['samples'][:2];m.update(status='partial',reason='one corner absent',impact='limited range',next_action='measure missing')
            result=CHECKER.validate(b,tmp)
            self.assertEqual(result['summaries'][0]['typ']['case_id'],'case_typ')
    def test_new_bundle_requires_calculation_lineage(self):
        with tempfile.TemporaryDirectory() as tmp:
            b=make_bundle(Path(tmp));b['manifest']['schema_version']='0.2'
            self.assertIn('calculation lacks input lineage',str(CHECKER.validate(b,tmp)['errors']))
            b['manifest']['artifacts'][2]['input_refs']=['a_net']
            self.assertTrue(CHECKER.validate(b,tmp)['valid'])
    def test_derived_invalid_run_propagation(self):
        with tempfile.TemporaryDirectory() as tmp:
            b=make_bundle(Path(tmp));b['manifest']['artifacts'][2]['input_refs']=['a_run']
            b['manifest']['runs'][0].update(status='excluded',reason='bad operating point')
            errors=str(CHECKER.validate(b,tmp)['errors'])
            self.assertIn('DERIVED sample inherits invalid run',errors)
            self.assertIn('DERIVED claim inherits invalid run',errors)
    def test_model_body_not_admitted_as_metadata_in_new_bundle(self):
        with tempfile.TemporaryDirectory() as tmp:
            b=make_bundle(Path(tmp));b['manifest']['schema_version']='0.2'
            b['manifest']['artifacts'][2]['input_refs']=['a_net']
            b['manifest']['artifacts'][0]['role']='model_metadata'
            self.assertIn('model metadata is not extracted',str(CHECKER.validate(b,tmp)['errors']))
    def test_a4_render_has_bounded_pixels(self):
        self.assertEqual(r.pixel_size(595,842),(1240,1755))
        with self.assertRaises(ValueError):r.pixel_size(595,842,1250)
    def test_each_verification_subsection_needs_all_labels(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'ok.docx';data=f.example_payload()
            data['verification'].append(dict(data['verification'][0], id='second_test', title='Second synthetic check'))
            f.build(data,path)
            profile=json.loads((f.PACKAGE/'check-profile.json').read_text(encoding='utf8'))
            self.assertTrue(d.validate(f.PACKAGE/'template.docx',path,profile)['valid'])
            doc=Document(path)
            labels=[p for p in doc.paragraphs if p.text.startswith('仿真方法：')]
            labels[-1]._element.getparent().remove(labels[-1]._element)
            bad=Path(tmp)/'bad.docx';doc.save(bad)
            self.assertIn('required label missing in subsection',str(d.validate(f.PACKAGE/'template.docx',bad,profile)['errors']))
    def test_page_field_removal_fails(self):
        from docx import Document
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'doc.docx';f.build(f.example_payload(),path);doc=Document(path)
            for part in doc.part.package.parts:
                el=getattr(part,'element',None)
                if el is not None:
                    for field in el.iter(qn('w:instrText')):
                        if 'NUMPAGES' in (field.text or ''):field.text=' PAGE '
            bad=Path(tmp)/'bad.docx';doc.save(bad)
            p=json.loads((f.PACKAGE/'check-profile.json').read_text(encoding='utf8'))
            self.assertTrue(any('template field removed: NUMPAGES' in e for e in d.validate(f.PACKAGE/'template.docx',bad,p)['errors']))
    def test_known_condition_cannot_hide_all_numbers(self):
        data=f.example_payload();data['sections']['1.3']['condition_rows'][0]['source_kind']='DIRECT'
        errors=f.production_preflight(data,None,None,None,None)
        self.assertIn('known condition needs numeric',str(errors))
    def test_interface_column_mismatch_fails(self):
        data=f.example_payload();data['sections']['1.4.1']['rows'][0]['cells'][2]='5V domain'
        self.assertIn('interface cells disagree',str(f.production_preflight(data,None,None,None,None)))
    def test_reading_a_method_is_not_completion(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);p=root/'skills/analog-review/SKILL.md';p.parent.mkdir(parents=True);p.write_text('Synthetic method')
            card=t.prepare('REVIEWER','Check synthetic ports',root,root)
            self.assertFalse(t.validate(card)['valid'])
            output=root/'review.txt';output.write_text('Synthetic review of two ports')
            card['permissions']['write_roots']=[str(root)]
            card.update(status='completed',inputs=[t.record(p)],outputs=[t.record(output)],acceptance=['ports match'],
                agent=dict(id='reviewer',model='gpt-5.6-luna',effort='high',model_selection_supported=True,delegated=True),
                author_id='author',requires_independent=True,closure=dict(reviewer_id='reviewer',evidence_refs=[str(output.resolve())]))
            self.assertIn('method read is not method result',t.validate(card)['errors'])
            card['method_results']=[dict(skill='analog-review',observation='Synthetic ports match at line 1',output_refs=[str(output.resolve())])]
            self.assertTrue(t.validate(card)['valid'])
            card['agent']['id']='author'
            self.assertIn('independence missing',t.validate(card)['errors'])

if __name__=='__main__':unittest.main()
