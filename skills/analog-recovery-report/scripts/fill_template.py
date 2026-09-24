"""Fill the reviewed fm-analog-v1 package. Never infer a template at run time.

Requires python-docx. Outputs a draft plus trace map; Word/PDF review is a
separate mandatory finish step. This validator checks structure, not physics.
"""
import argparse
import copy
import hashlib
import json
import re
import importlib.util
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor, Pt, Inches
from docx.table import Table

PACKAGE = Path(__file__).resolve().parents[1] / 'assets/fm-analog-v1'
LABELS = ('仿真工具', '仿真路径', '仿真方法', '结果分析')
CONDITION_QUANTITIES = ('supply_voltage', 'temperature')
SOURCE_KINDS = ('DIRECT', 'DERIVED', 'SIMULATED', 'INFERRED', 'UNKNOWN')


def verification_rows(data):
    """One source for the template plan and the corresponding result sections."""
    return [{'cells': [c['category'], c['title'], c['status'], c['description']],
             'refs': c['refs']} for c in data['verification']]


def verification_groups(data):
    """Keep partial results with figures; consolidate only N items without figures."""
    results, gaps = [], []
    for case in data['verification']:
        (gaps if case['status'] == 'N' and not case.get('figures') else results).append(case)
    return results, gaps


def metric_display(metric, compact):
    """One scale per metric, never change evidence values or compound units."""
    unit = metric['unit']
    values = [abs(metric[k]['value']) for k in ('min', 'typ', 'max') if metric.get(k)]
    magnitude = max(values, default=0)
    if compact and unit in ('V', 'A', 'W', 's') and magnitude:
        for scale, prefix in ((1e9, 'G'), (1e6, 'M'), (1e3, 'k'), (1, ''),
                              (1e-3, 'm'), (1e-6, 'µ'), (1e-9, 'n'), (1e-12, 'p')):
            if scale <= magnitude < scale * 1000:
                return scale, prefix + unit
    return 1, unit


def group_interfaces(rows):
    # Preserve the declared state; do not infer pin roles from their names.
    return sorted(rows, key=lambda r: (r['cells'][1], r['cells'][0]))


def compact_conditions(metric):
    """Factor exact common tokens; retain full source cases in the trace sidecar."""
    selected=[(k,metric.get(k)) for k in ('min','typ','max') if metric.get(k)]
    parts={s['case_id']:[t.strip() for t in re.split('[，,；;]',metric['conditions'][s['case_id']]['conditions']) if t.strip()] for _,s in selected}
    common=[t for t in next(iter(parts.values()),[]) if all(t in v for v in parts.values())]
    lines=['；'.join(common)] if common else []
    seen=set()
    for _,s in selected:
        cid=s['case_id']
        if cid in seen:continue
        seen.add(cid)
        labels='/'.join({'min':'最小','typ':'典型','max':'最大'}[k] for k,v in selected if v['case_id']==cid)
        diff=[t for t in parts[cid] if t not in common]
        if diff:lines.append(labels+'：'+'；'.join(diff))
    return '\n'.join(lines)


def profile():
    return json.loads((PACKAGE/'profile.json').read_text(encoding='utf-8'))


def shape_errors(value, rule, path='$'):
    """Validate the closed schema vocabulary used by this bundled content schema.

    Not a general JSON Schema implementation; no refs/unions/remote resolution.
    """
    errors=[]
    types={'object':dict,'array':list,'string':str}
    if 'type' in rule and not isinstance(value,types[rule['type']]):return [path+': wrong type']
    if 'enum' in rule and value not in rule['enum']:errors.append(path+': invalid enum')
    if isinstance(value,str) and len(value.strip())<rule.get('minLength',0):errors.append(path+': empty string')
    if isinstance(value,dict):
        props=rule.get('properties',{})
        for key in rule.get('required',[]):
            if key not in value:errors.append(path+': missing '+key)
        if rule.get('additionalProperties') is False and set(value)-set(props):errors.append(path+': unknown fields')
        for key,item in value.items():
            if key in props:errors.extend(shape_errors(item,props[key],path+'.'+key))
    if isinstance(value,list):
        if len(value)<rule.get('minItems',0):errors.append(path+': too few items')
        if rule.get('uniqueItems') and len({json.dumps(v,sort_keys=True) for v in value})!=len(value):errors.append(path+': duplicate items')
        for i,item in enumerate(value):errors.extend(shape_errors(item,rule.get('items',{}),path+'.'+str(i)))
    return errors


def validate(data):
    errors=[]
    try:
        schema=json.loads((PACKAGE/'content.schema.json').read_text(encoding='utf-8'))
        syntax=shape_errors(data,schema)
        if syntax:return syntax
        p=profile(); sections=data['sections']; ids=set(data['evidence_ids'])
        if data['template_id'] != p['id']: errors.append('template_id mismatch')
        if not all(isinstance(data[k],str) and data[k].strip() for k in ('title','project_id','project_name','date','revision','author')):
            errors.append('missing cover metadata')
        expected=set(p['sections'])
        if set(sections)!=expected: errors.append('section keys mismatch: '+str(set(sections)^expected))
        def refs(item):
            if not item.get('refs') or not set(item['refs'])<=ids: errors.append('missing/unknown evidence refs')
        for key,s in sections.items():
            if key not in expected: continue
            if s.get('status') not in ('supported','partial','unverified','not_applicable'): errors.append('invalid section status '+key)
            if s.get('status')!='supported' and not s.get('reason'): errors.append('missing gap reason '+key)
            if s.get('status')=='supported' and key in ('1.3','1.4.1','1.7.2') and not s.get('rows'):
                errors.append('supported section missing required table '+key)
            if s.get('status')=='supported' and key=='1.5' and not s.get('figures'):
                errors.append('supported architecture missing diagram')
            for item in s.get('paragraphs',[]):
                refs(item)
                if not isinstance(item.get('text'),str) or not item['text'].strip(): errors.append('empty paragraph '+key)
                elif re.match(r'^\s*(?:\d+[.)、]|[*★])',item['text']): errors.append('manual numbering '+key)
            for item in s.get('figures',[]):
                refs(item)
                if item.get('kind') not in ('concept','connectivity','simulation','analytical'): errors.append('invalid figure kind')
                if not item.get('caption') or not Path(item.get('path','')).is_file(): errors.append('missing figure file/caption')
            for row in s.get('rows',[]):
                refs(row)
                if key not in p['tables'] or len(row.get('cells',[]))!=p['tables'][key]['columns']: errors.append('table width '+key)
            if key == '1.3':
                condition_rows = s.get('condition_rows')
                if condition_rows is None:
                    errors.append('1.3 missing condition_rows; add supply_voltage and temperature rows (UNKNOWN allowed)')
                else:
                    seen = set()
                    for row in condition_rows:
                        quantity = row.get('quantity')
                        if quantity not in CONDITION_QUANTITIES:
                            errors.append('1.3 invalid condition quantity: '+str(quantity))
                        elif quantity in seen:
                            errors.append('1.3 duplicate condition quantity: '+quantity)
                        else:
                            seen.add(quantity)
                        if len(row.get('cells',[])) != p['tables']['1.3']['columns']:
                            errors.append('1.3 condition row width')
                        if row.get('source_kind') not in SOURCE_KINDS:
                            errors.append('1.3 condition row source_kind must be DIRECT/DERIVED/SIMULATED/INFERRED/UNKNOWN')
                        refs(row)
                    missing = set(CONDITION_QUANTITIES)-seen
                    if missing: errors.append('1.3 missing condition rows: '+', '.join(sorted(missing)))
            is_parent=any(other.startswith(key+'.') for other in expected)
            if not (s.get('paragraphs') or s.get('rows') or s.get('figures') or s.get('reason') or key=='1.8' or is_parent):
                errors.append('empty section '+key)
        covered=[port for r in sections.get('1.4.1',{}).get('rows',[]) for port in r.get('ports',[])]
        if len(covered)!=len(set(covered)) or set(covered)!=set(data['expected_ports']): errors.append('interface port coverage mismatch')
        if not data.get('verification'): errors.append('verification needs explicit result or gap')
        if sections['1.8'].get('figures'):
            errors.append('put 1.8 figures inside the corresponding verification item')
        if sections['1.6'].get('rows') and sections['1.6']['rows'] != verification_rows(data):
            errors.append('1.6 rows must equal generated verification plan; omit rows to generate automatically')
        model_names = [r['cells'][0].strip() for r in sections['1.7.2'].get('rows', [])]
        if len(model_names) != len(set(model_names)):
            errors.append('duplicate model in 1.7.2; put instance sizing in 1.7.1')
        metric_ids = {r['metric_ref'] for r in sections['1.3'].get('rows', []) if r.get('metric_ref')}
        covered_metrics = {m for c in data['verification'] for m in c['metric_refs']}
        if metric_ids - covered_metrics: errors.append('uncovered metric in verification: '+str(sorted(metric_ids-covered_metrics)))
        if covered_metrics - ids: errors.append('verification metric_refs contain unknown evidence IDs')
        case_ids = [c['id'] for c in data['verification']]
        if len(case_ids) != len(set(case_ids)): errors.append('duplicate verification id')
        for case in data.get('verification',[]):
            refs(case)
            for fig in case.get('figures', []):
                refs(fig)
                if not Path(fig['path']).is_file(): errors.append('missing verification figure file')
            if not case.get('title') or set(case.get('labels',{}))!=set(LABELS): errors.append('verification title/labels incomplete')
            if any(not isinstance(v,str) or not v.strip() for v in case.get('labels',{}).values()): errors.append('empty verification label')
    except (KeyError,TypeError,ValueError) as exc:
        errors.append('malformed content: '+str(exc))
    return errors


def bind_bundle(data, bundle_path, source_root):
    """Recheck the actual bundle; derive metric cells instead of trusting prose.

    This complements the independent electrical gate, not a replacement for it.
    """
    script=Path(__file__).resolve().parents[2]/'analog-recovery-evidence/scripts/check_bundle.py'
    spec=importlib.util.spec_from_file_location('recovery_bundle_for_report',script)
    mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    bundle={n:mod.read(Path(bundle_path)/(n+'.json')) for n in mod.FILES}
    result=mod.validate(bundle,source_root)
    if not result['valid']:raise ValueError('evidence bundle invalid: '+str(result['errors']))
    bound=copy.deepcopy(data)
    actual_ids={v['id'] for key,seq in [('design_evidence','claims'),('design_evidence','modes'),('metrics','metrics'),('manifest','artifacts'),('manifest','runs')] for v in bundle[key][seq]}
    if not set(data['evidence_ids'])<=actual_ids:raise ValueError('content invents evidence IDs')
    bound['evidence_ids']=sorted(actual_ids)
    bound['expected_ports']=[v['name'] for v in bundle['design_evidence']['ports']]
    metrics={m['id']:m for m in result['summaries']}
    definitions={m['id']:m for m in bundle['metrics']['metrics']}
    for row in bound['sections']['1.3'].get('rows',[]):
        ref=row.get('metric_ref')
        if ref not in metrics or ref not in row['refs']:raise ValueError('metric row needs bundle metric_ref and refs')
        metric=metrics[ref]
        scale, display_unit = metric_display(metric, bound.get('compact_metrics'))
        def display_value(item):
            return format(item['value'] / scale, '.5g' if bound.get('compact_metrics') else '.6g')
        row['condition_cases'] = metric.get('conditions', {})
        def value(k):
            item=metric.get(k)
            return display_value(item) if item is not None else ('未指定典型工况' if k=='typ' and metric.get('min') and definitions[ref]['typical_case_id'] is None else '未测')
        row['cells'][4:8]=['—' if metric['range_interpretation']=='single_point_only' else value('min'),value('typ'),'—' if metric['range_interpretation']=='single_point_only' else value('max'),display_unit]
        # Author prose cannot override conditions of the actual selected samples.
        values=[]
        scope=str(metric.get('scope',''))
        shared={t.strip() for t in re.split('[，,；;]',scope) if t.strip()}
        for label in ('min','typ','max'):
            sample=metric.get(label)
            if sample:
                case=metric['conditions'][sample['case_id']]
                remaining=[t.strip() for t in re.split('[，,；;]',case['conditions']) if t.strip() and t.strip() not in shared]
                values.append(label+' ['+sample['case_id']+'] '+('，'.join(remaining) or '同上述条件'))
        row['cells'][3] = scope+'\n'+'\n'.join(values)
        if bound.get('compact_metrics'):
            row['cells'][3]=compact_conditions(metric)
            if row['cells'][5]=='未指定典型工况':row['cells'][5]='—';row['cells'][3]+='；典型条件未定义'
        if metric['range_interpretation']=='single_point_only':
            sample = metric.get('min') or metric.get('max') or metric.get('typ')
            derived = definitions[ref]['kind'] == 'DERIVED'
            row['cells'][3]+='；单个派生结果，不代表范围' if derived else '；仅单工况实测，不代表范围'
            if sample and metric.get('typ') is None:
                row['cells'][2]+='\n本次结果：'+display_value(sample)+' '+display_unit
                if not bound.get('compact_metrics'):
                    row['cells'][3]+= ('；派生值 ' if derived else '；已测 ')+display_value(sample)+' '+display_unit
    return bound,result


def check_electrical(bundle_path, packet_path):
    """Require the production electrical gate for metrics backed by simulation.

    Dimensional validation is not numerical formula evaluation. Bundle values
    and raw values still require extraction/review in the evidence stage.
    """
    metrics=json.loads((Path(bundle_path)/'metrics.json').read_text(encoding='utf-8'))['metrics']
    measured=[m for m in metrics if any(s.get('run_ref') for s in m.get('samples',[]))]
    if not measured:return None
    if not packet_path:raise ValueError('simulation metrics require --measurement-packet')
    script=Path(__file__).resolve().parents[2]/'analog-recovery-evidence/scripts/check_measurements.py'
    spec=importlib.util.spec_from_file_location('report_electrical_gate',script)
    mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    path=Path(packet_path).resolve();packet=mod.load_packet(path,synthetic=False)
    result=mod.validate(packet,base_dir=path.parent)
    if not result['valid']:raise ValueError('electrical gate failed: '+str(result['errors']))
    declared={m['id']:m for m in packet['measurements']};runs={r['id']:r for r in packet['runs']}
    result['metric_units_requiring_independent_review']=[]
    for metric in measured:
        if re.fullmatch(r'[fpnumkMGT]?[VAW]',metric['unit']):
            if metric['id'] not in declared or declared[metric['id']]['unit']!=metric['unit']:
                raise ValueError('electrical packet missing metric or unit mismatch: '+metric['id'])
        else:
            result['metric_units_requiring_independent_review'].append({'id':metric['id'],'unit':metric['unit']})
        for sample in metric['samples']:
            rid=sample.get('run_ref')
            if rid and (rid not in runs or runs[rid]['status'] not in ('valid','qualified')):
                raise ValueError('metric uses unaccepted electrical run: '+str(rid))
    return result


INTERFACE_FIELDS=('name','direction','hysteresis','fan_in_out','description')

def production_preflight(data,content_path,bundle_path,source_root,quality_path):
    """Bind the reviewed snapshot and reject known template column misuse.

    This verifies declarations and file identities, not arbitrary prose truth.
    """
    errors=[]
    for row in data['sections']['1.4.1'].get('rows',[]):
        fields=row.get('interface')
        if not fields or set(fields)!=set(INTERFACE_FIELDS):
            errors.append('interface requires named name/direction/hysteresis/fan_in_out/description');continue
        if row['cells']!=[fields[k] for k in INTERFACE_FIELDS]:errors.append('interface cells disagree with named fields')
    for row in data['sections']['1.3'].get('condition_rows',[]):
        if row['source_kind']!='UNKNOWN' and all(str(v).strip() in ('—','-','未测','UNKNOWN','') for v in row['cells'][4:7]):
            errors.append('known condition needs numeric min/typ/max slots: '+row['quantity'])
    if not quality_path:return errors+['production fill requires --quality-packet (bounded engineering checks)']
    script=Path(__file__).resolve().parents[2]/'analog-recovery-evidence/scripts/check_quality.py'
    spec=importlib.util.spec_from_file_location('report_quality',script);mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    packet=json.loads(Path(quality_path).read_text(encoding='utf-8-sig'))
    result=mod.validate(packet,source_root)
    errors.extend(str(e) for e in result['errors'])
    bundle={n:json.loads((Path(bundle_path)/(n+'.json')).read_text(encoding='utf-8-sig')) for n in ('manifest','design_evidence','metrics')}
    if bundle['manifest']['schema_version']!='0.2':errors.append('production requires a new 0.2 bundle with calculation lineage; preserve legacy bundle')
    declared_runs={r['id']:r for r in bundle['manifest']['runs']}
    declared_claims={c['id']:c for c in bundle['design_evidence']['claims']}
    declared_metrics={m['id']:m for m in bundle['metrics']['metrics'] if m['samples']}
    for run in packet.get('runs',[]):
        if run['id'] not in declared_runs or run['status']!=declared_runs[run['id']]['status']:errors.append('quality run differs from evidence: '+run['id'])
    for claim in packet.get('claims',[]):
        original=declared_claims.get(claim['id'])
        if not original:errors.append('quality invents claim: '+claim['id'])
        elif original['kind']=='UNKNOWN' and claim['status']!='unknown':errors.append('quality promotes UNKNOWN: '+claim['id'])
        elif not set(original['run_refs'])<=set(claim['run_refs']):errors.append('quality drops claim run: '+claim['id'])
    qmetrics={m['id']:m for m in packet.get('metrics',[])}
    if set(qmetrics)!=set(declared_metrics):errors.append('quality must cover all measured metric conditions')
    for mid,m in declared_metrics.items():
        if mid in qmetrics:
            qm=qmetrics[mid]
            if set(qm['case_ids'])!={s['case_id'] for s in m['samples']} or qm['typical_case_id']!=m['typical_case_id']:errors.append('quality metric differs from evidence: '+mid)
    report_texts=[p['text'] for s in data['sections'].values() for p in s.get('paragraphs',[])]+[v for c in data['verification'] for v in c['labels'].values()]+[g['caption'] for s in data['sections'].values() for g in s.get('figures',[])]+[g['caption'] for c in data['verification'] for g in c.get('figures',[])]
    for use in packet.get('uses',[]):
        if use['text'] not in report_texts:errors.append('quality use is not current report text: '+use['claim_ref'])
    frozen=packet.get('report_snapshot',{})
    if frozen.get('content_sha256')!=mod.digest(content_path):errors.append('quality review content snapshot missing/changed')
    for name in ('manifest','report_requirements','design_evidence','metrics','coverage','review'):
        if frozen.get('bundle_sha256',{}).get(name)!=mod.digest(Path(bundle_path)/(name+'.json')):errors.append('quality review bundle snapshot missing/changed: '+name)
    return errors


def clean_run(run):
    run.font.color.rgb=RGBColor(0,0,0)
    run.font.italic=False
    run.font.underline=False


def normalize_docpr_ids(doc):
    """Give every drawing property a package-global unique ID.

    Word validates ``wp:docPr/@id`` across the main story, headers, and
    footers.  After the fixed-template body is truncated, python-docx may
    restart inserted-picture IDs at 1 even though the retained header/footer
    parts already use that ID, which Word reports as a corrupt document.
    """
    next_id = 1
    for part in doc.part.package.parts:
        element = getattr(part, 'element', None)
        if element is None:
            continue
        for node in element.iter(qn('wp:docPr')):
            node.set('id', str(next_id))
            next_id += 1


def paragraph(doc,text,prototype=None):
    p=doc.add_paragraph()
    if prototype is not None and prototype._p.pPr is not None:
        p._p.insert(0,copy.deepcopy(prototype._p.pPr))
    else:
        p.style=doc.styles['正文首行缩进']
        # Template style can inherit instructional blue/italic. Direct overrides
        # are deliberate; paragraph styles, numbering and fonts stay native.
    p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text);clean_run(r)
    return p


def fill_table(doc,prototype,rows,header_rows,columns,column_weights=None):
    el=copy.deepcopy(prototype._tbl)
    for row in list(el.findall(qn('w:tr')))[header_rows:]: el.remove(row)
    doc._body._body.insert(len(doc._body._body)-1,el)
    table=Table(el,doc._body)
    widths=list(el.find(qn('w:tblGrid')))
    if columns==8 or column_weights is not None:
        total=sum(int(w.get(qn('w:w'))) for w in widths)
        weights=column_weights or [0.8,1.7,3.1,4.2,1.35,1.35,1.35,1.0]
        sizes=[round(total*w/sum(weights)) for w in weights]
        sizes[-1]+=total-sum(sizes)
        for col,size in zip(widths,sizes):col.set(qn('w:w'),str(size))
        for tr in el.findall(qn('w:tr')):
            index=0
            for tc in tr.findall(qn('w:tc')):
                pr=tc.find(qn('w:tcPr'));span=pr.find(qn('w:gridSpan'));count=int(span.get(qn('w:val'))) if span is not None else 1
                pr.find(qn('w:tcW')).set(qn('w:w'),str(sum(sizes[index:index+count])));index+=count
    for row in rows:
        tr=OxmlElement('w:tr')
        for i,value in enumerate(row['cells']):
            tc=OxmlElement('w:tc'); pr=OxmlElement('w:tcPr'); w=OxmlElement('w:tcW')
            w.set(qn('w:w'),widths[i].get(qn('w:w')));w.set(qn('w:type'),'dxa');pr.append(w);tc.append(pr)
            if columns==8 and i in (4,5,6,7):pr.append(OxmlElement('w:noWrap'))
            p=OxmlElement('w:p');r=OxmlElement('w:r')
            for j,line in enumerate(str(value).split('\n')):
                if j:r.append(OxmlElement('w:br'))
                t=OxmlElement('w:t');t.text=line;r.append(t)
            p.append(r);tc.append(p);tr.append(tc)
        el.append(tr)
    # Keep the template grid and headers. New body rows need an explicit native
    # paragraph style; otherwise they inherit dense bold/default formatting.
    for row in table.rows[header_rows:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.style=doc.styles['Normal']
                p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.space_before=Pt(3)
                for r in p.runs:
                    r.bold=False;r.font.size=Pt(10)
    for row in table.rows:
        # Keep one engineering record intact across pages (e.g. a unit must
        # not be separated from its number at a page boundary).
        pr = row._tr.get_or_add_trPr()
        if pr.find(qn('w:cantSplit')) is None:
            pr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs: clean_run(r)
    for row in table.rows[:header_rows]:
        pr=row._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');pr.append(repeat)
    return table


def build(data,out,template=None):
    cfg=profile(); template=Path(template or PACKAGE/'template.docx'); out=Path(out)
    if hashlib.sha256(template.read_bytes()).hexdigest()!=cfg['sha256']: raise ValueError('template changed; explicit adaptation required')
    problems=validate(data)
    if problems: raise ValueError('; '.join(problems))
    data=copy.deepcopy(data)
    data['sections']['1.6']['rows']=verification_rows(data)
    data['sections']['1.4.1']['rows']=group_interfaces(data['sections']['1.4.1'].get('rows', []))
    if out.exists(): raise FileExistsError(out)
    src=Document(template);doc=Document(template)
    body=doc._body._body
    # Keep the exact front matter through its directory heading, including the
    # revision table. Remove old TOC cached entries, examples and second module.
    cut=doc.paragraphs[19]._p
    removing=False
    for el in list(body):
        if el is cut: removing=True
        if removing and el.tag!=qn('w:sectPr'):body.remove(el)
    for idx,text in ((6,data['title']),(10,'项目编号：'+data['project_id']),(11,'项目名称：'+data['project_name'])):
        p=doc.paragraphs[idx];p.text=text
        for r in p.runs:clean_run(r)
    rev=doc.tables[0]
    for row in rev.rows[1:]:
        for cell in row.cells:cell.text=''
    for c,v in zip(rev.rows[1].cells,(data['date'],data['revision'],data.get('change','证据重建报告'),data['author'])):c.text=v
    toc=doc.add_paragraph();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u');toc._p.append(fld)
    heading=copy.deepcopy(src.paragraphs[78]._p);body.insert(len(body)-1,heading)
    doc.paragraphs[-1].text=data['title']
    doc.paragraphs[-1].paragraph_format.page_break_before=True
    trace=[]
    def figure(item, section):
        doc.add_picture(item['path'],width=Inches(5.7))
        paragraph(doc,item['caption']+'（'+{'concept':'概念示意','connectivity':'连接图','simulation':'实际仿真数据','analytical':'解析计算，非仿真'}[item['kind']]+'）')
        trace.append({'section':section,**item})
    def content(key):
        s=data['sections'][key]
        if s.get('reason'):
            trace.append({'section':key,'kind':'section_reason','reason':s['reason']})
        for item in s.get('paragraphs',[]):
            paragraph(doc,item['text'],src.paragraphs[84] if key in ('1.1','1.2') else None)
            trace.append({'section':key,'kind':'paragraph',**item})
        if s.get('rows') or (key == '1.3' and s.get('condition_rows')):
            table_rows = list(s.get('condition_rows', [])) + list(s['rows']) if key == '1.3' else s['rows']
            t=cfg['tables'][key];fill_table(doc,src.tables[t['index']],table_rows,t['header_rows'],t['columns'])
            if key == '1.3':
                trace.extend({'section':key,'kind':'condition_row',**r} for r in s.get('condition_rows', []))
            trace.extend({'section':key,'kind':'table_row',**r} for r in s['rows'])
        for item in s.get('figures',[]):
            figure(item, key)
    for key,s in cfg['sections'].items():
        h=copy.deepcopy(src.paragraphs[s['paragraph_index']]._p);body.insert(len(body)-1,h)
        content(key)
        if key=='1.8':
            result_cases, gap_cases = verification_groups(data)
            for case in result_cases:
                h=copy.deepcopy(src.paragraphs[138]._p);body.insert(len(body)-1,h);doc.paragraphs[-1].text=case['title']
                for label in LABELS:
                    if label=='结果分析':
                        for item in case.get('figures',[]): figure(item,key+'.'+case['id'])
                    paragraph(doc,label+'：'+case['labels'][label],src.paragraphs[143])
                trace.append({'section':key,'kind':'verification',**case})
            if gap_cases:
                h=copy.deepcopy(src.paragraphs[138]._p);body.insert(len(body)-1,h)
                doc.paragraphs[-1].text='未完成项与后续条件'
                # Retain the native four-label contract once, not four empty
                # paragraphs for every unperformed test. Full records stay in trace.
                for label, text in zip(LABELS, (
                    '以下项目尚无本次接受的完整验证结果；已尝试的工具和日志见证据索引。',
                    '按下表项目定位1.6验证计划及证据索引中的原始记录。',
                    '各项补充条件与下一步见下表；本节不代表新增运行。',
                    '缺口及其影响如下，不将未完成项记为验证通过。')):
                    paragraph(doc,label+'：'+text,src.paragraphs[143])
                rows=[{'cells':[c['title'],c['description'],c['labels']['仿真方法'],c['labels']['结果分析']],
                       'refs':c['refs']} for c in gap_cases]
                t=cfg['tables']['1.6']
                table=fill_table(doc,src.tables[t['index']],rows,t['header_rows'],t['columns'],[2,3,4,3])
                for cell, title in zip(table.rows[0].cells, ('项目','当前状态/原因','补充条件/下一步','影响/结论限制')):
                    cell.text=title
                    for run in cell.paragraphs[0].runs: clean_run(run)
                trace.extend({'section':key,'kind':'verification_gap',**c} for c in gap_cases)
    for section in doc.sections:
        for part in (section.header,section.footer,section.first_page_header,section.first_page_footer):
            for p in part.paragraphs:
                for r in p.runs:
                    if '<FMXXX>' in r.text:r.text=r.text.replace('<FMXXX>',data['project_id'])
    normalize_docpr_ids(doc)
    settings=doc.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
    out.parent.mkdir(parents=True,exist_ok=True);doc.save(out)
    out.with_suffix('.trace.json').write_text(json.dumps({'template':cfg['id'],'template_sha256':cfg['sha256'],'status':'draft_requires_field_update_and_render_review','blocks':trace},ensure_ascii=False,indent=2),encoding='utf-8')
    return out


def audit(path):
    d=Document(path);errors=[];active=False
    for p in d.paragraphs:
        if p.style.name=='Heading 1':active=True
        if not active:continue
        if re.search(r'此处描述|此章节可指导|<FMXXX>|Mismatch=1',p.text):errors.append('template residue: '+p.text[:70])
        if p.style.name.startswith('Heading'):continue
        for r in p.runs:
            if r.text and (r.italic or str(r.font.color.rgb) not in ('None','000000')):errors.append('body instruction formatting: '+r.text[:50])
        if p.text.startswith(tuple(x+'：' for x in LABELS)) and p._p.find('.//'+qn('w:numPr')) is not None:errors.append('numbered fixed label')
    return {'errors':errors,'scope':'body direct formatting/residue; render review still required'}


def example_payload():
    p=profile();sections={k:{'status':'unverified','reason':'独立合成测试，不代表任何产品电路。','paragraphs':[{'text':'此段用于测试结构化内容填充与证据引用。','refs':['synthetic']}]} for k in p['sections']}
    sections['1.3']['condition_rows']=[
        {'quantity':'supply_voltage','cells':['工作条件','VDD','供电电压','UNKNOWN','—','—','—','V'],'refs':['synthetic'],'source_kind':'UNKNOWN'},
        {'quantity':'temperature','cells':['工作条件','温度','环境温度','UNKNOWN','—','—','—','°C'],'refs':['synthetic'],'source_kind':'UNKNOWN'},
    ]
    sections['1.4.1']['rows']=[{'cells':['OUT','输出','未测','未测','合成端口，仅检验覆盖映射'],'ports':['OUT'],'refs':['synthetic']}]
    sections['1.4.1']['rows'][0]['interface']=dict(zip(INTERFACE_FIELDS,sections['1.4.1']['rows'][0]['cells']))
    sections['1.3']['rows']=[{'cells':['测试项','VOUT','合成填充','未执行仿真','—','—','—','V'],'refs':['synthetic']}]
    return {'template_id':p['id'],'title':'独立合成模板验收样本','project_id':'SYNTHETIC','project_name':'模板生成器测试','date':'2026-09-19','revision':'test','author':'自动化测试','evidence_ids':['synthetic'],'expected_ports':['OUT'],'sections':sections,'verification':[{'id':'synthetic_test','category':'功能','status':'N','description':'未运行电气验证；仅检验模板填充','metric_refs':[],'title':'合成流程检查','labels':dict(zip(LABELS,['未运行仿真','不适用','仅验证模板填充','无电气结论'])),'refs':['synthetic']}]}


if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--content');ap.add_argument('--out');ap.add_argument('--example');ap.add_argument('--requirements');ap.add_argument('--audit');ap.add_argument('--bundle');ap.add_argument('--source-root');ap.add_argument('--measurement-packet');ap.add_argument('--quality-packet');ap.add_argument('--check',action='store_true');ap.add_argument('--synthetic',action='store_true');a=ap.parse_args()
    if a.example:Path(a.example).write_text(json.dumps(example_payload(),ensure_ascii=False,indent=2),encoding='utf-8')
    elif a.requirements:
        req=json.loads((PACKAGE/'requirements.json').read_text(encoding='utf-8'))
        output={'requirements':[{'id':r['id'],'section':r['section'],'request':r['task']+'；验收：'+r['acceptance'],'required':True} for r in req['requirements']]}
        Path(a.requirements).write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
    elif a.audit:print(json.dumps(audit(a.audit),ensure_ascii=False,indent=2))
    else:
        data=json.loads(Path(a.content).read_text(encoding='utf-8'))
        errors=validate(data)
        bundle_status={}
        if not a.synthetic:
            if not a.bundle or not a.source_root:ap.error('production fill needs --bundle and --source-root; only independent fixtures use --synthetic')
            if not errors:errors.extend(production_preflight(data,a.content,a.bundle,a.source_root,a.quality_packet))
            for stage in ('bundle','electrical'):
                try:
                    if stage=='bundle':
                        data,summary=bind_bundle(data,a.bundle,a.source_root)
                        bundle_status={'bundle_delivery_status':summary['delivery_status'],'key_claim_review':summary['key_claim_review']}
                    else:electrical=check_electrical(a.bundle,a.measurement_packet)
                except (ValueError,KeyError,TypeError,OSError) as exc:errors.append(stage+': '+str(exc))
        elif not data['project_id'].startswith('SYNTHETIC'):ap.error('--synthetic requires a SYNTHETIC project_id')
        if errors or a.check:
            print(json.dumps({'valid':not errors,'errors':errors,**bundle_status},ensure_ascii=False,indent=2));raise SystemExit(bool(errors))
        if not a.out:ap.error('--out required unless --check')
        build(data,a.out)
        if not a.synthetic:
            Path(a.out).with_suffix('.electrical.json').write_text(json.dumps(electrical,ensure_ascii=False,indent=2),encoding='utf-8')
        receipt={'content_sha256':hashlib.sha256(Path(a.content).read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256(Path(a.out).read_bytes()).hexdigest(),'filler_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'template_sha256':profile()['sha256']}
        receipt.update(bundle_status)
        if a.quality_packet:receipt['quality_sha256']=hashlib.sha256(Path(a.quality_packet).read_bytes()).hexdigest()
        Path(a.out).with_suffix('.build.json').write_text(json.dumps(receipt,indent=2),encoding='utf-8')
        print(json.dumps(audit(a.out),ensure_ascii=False))
