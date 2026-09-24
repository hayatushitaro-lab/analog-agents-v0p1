"""Targeted offline regressions for recovery closeout and readable report output."""
import copy
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image

from test_skill_handoff import make_bundle

ROOT = Path(__file__).resolve().parents[1]


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, ROOT / path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


f = load('report_closeout_filler', 'skills/analog-recovery-report/scripts/fill_template.py')
w = load('report_closeout_workflow', 'skills/analog-recovery/scripts/resolve_workflow.py')
c = load('report_closeout_checker', 'skills/analog-recovery-report/scripts/check_docx.py')


class CloseoutReportingTests(unittest.TestCase):
    def test_plan_does_not_claim_evolve_execution_or_expand_authority(self):
        for tier in ('lite', 'standard', 'intensive', 'exhaustive'):
            plan = w.resolve(dict(workflow_effort=tier, knowledge_write_policy='read_only', single_agent=True))
            execution = plan['closeout']['evolution_execution']
            self.assertEqual(execution['required'], tier != 'lite')
            self.assertEqual(execution['status'], 'not_required_by_effort' if tier == 'lite' else 'pending')
            self.assertEqual(execution['engine_execution'], 'not_run')
            self.assertTrue((ROOT / execution['skill_path']).is_file())
            self.assertEqual(plan['coverage_closeout']['status'], 'pending')
            self.assertFalse(plan['execution']['remote_execution_allowed'])
            self.assertFalse(plan['dispatch']['starts_agents'])
            self.assertFalse(plan['review_transport']['authorized'])
            self.assertEqual(plan['closeout']['write_policy'], 'read_only')
            self.assertFalse(plan['closeout']['new_project_convergence'])

    def test_one_scale_for_signed_extrema_and_no_compound_unit_conversion(self):
        metric = dict(unit='A', min={'value': -2e-6}, typ={'value': 0}, max={'value': 8e-6})
        before = copy.deepcopy(metric)
        self.assertEqual((1e-6, 'µA'), f.metric_display(metric, True))
        self.assertEqual((1, 'A'), f.metric_display(metric, False))
        self.assertEqual(metric, before)
        metric['unit'] = 'ppm/°C'
        self.assertEqual((1, 'ppm/°C'), f.metric_display(metric, True))
        metric = dict(unit='W', min={'value': 0}, max={'value': 0})
        self.assertEqual((1, 'W'), f.metric_display(metric, True))

    def test_derived_scalar_visible_without_inventing_typ_or_single_temperature(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            bundle = make_bundle(root)
            metric = bundle['metrics']['metrics'][0]
            metric.update(kind='DERIVED', unit='ppm/°C', typical_case_id=None)
            metric['samples'] = [dict(case_id='three_temperatures', value=35.7, source_ref='a_calc', locator='fixture:1', run_ref=None)]
            metric['expected_cases'] = [dict(id='three_temperatures', conditions='TT; -40/27/125°C; three input samples')]
            for name, value in bundle.items():
                (root / (name + '.json')).write_text(json.dumps(value), encoding='utf-8')
            original = (root / 'metrics.json').read_bytes()
            payload = f.example_payload()
            payload.update(compact_metrics=True, evidence_ids=[metric['id']])
            row = payload['sections']['1.3']['rows'][0]
            row.update(metric_ref=metric['id'], refs=[metric['id']])
            bound, result = f.bind_bundle(payload, root, root)
            cells = bound['sections']['1.3']['rows'][0]['cells']
            self.assertIn('35.7 ppm/°C', cells[2])
            self.assertNotIn('仅单点', cells[3])
            self.assertIn('-40/27/125°C', cells[3])
            self.assertEqual(['—', '—', '—'], cells[4:7])
            self.assertIsNone(result['summaries'][0]['typ'])
            self.assertEqual(original, (root / 'metrics.json').read_bytes())

    def test_gap_consolidation_keeps_plan_partial_figure_and_trace(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            figure = root / 'result.png'
            Image.new('RGB', (160, 80), 'white').save(figure)
            payload = f.example_payload()
            base = payload['verification'][0]
            payload['verification'] = []
            for i in range(3):
                case = copy.deepcopy(base)
                case.update(id=f'case{i}', title=f'TEST_CASE_{i}', description=f'REASON_{i}')
                case['labels']['仿真方法'] = f'NEXT_ACTION_{i}'
                case['labels']['结果分析'] = f'IMPACT_{i}'
                payload['verification'].append(case)
            # N with genuine partial result still has its own subsection/figure.
            payload['verification'][0]['figures'] = [dict(path=str(figure), caption='PARTIAL_RESULT', kind='analytical', refs=['synthetic'])]
            results, gaps = f.verification_groups(payload)
            self.assertEqual(['case0'], [x['id'] for x in results])
            self.assertEqual(['case1', 'case2'], [x['id'] for x in gaps])
            self.assertEqual(3, len(f.verification_rows(payload)))
            out = root / 'report.docx'
            f.build(payload, out)
            doc = Document(out)
            text = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('TEST_CASE_0', text)
            self.assertNotIn('TEST_CASE_1', text)
            table_text = '\n'.join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
            for i in range(3):
                self.assertIn(f'TEST_CASE_{i}', table_text)
            for i in (1, 2):
                self.assertIn(f'NEXT_ACTION_{i}', table_text)
                self.assertIn(f'IMPACT_{i}', table_text)
            trace = json.loads(out.with_suffix('.trace.json').read_text(encoding='utf-8'))
            records = [x for x in trace['blocks'] if x.get('kind') in ('verification', 'verification_gap')]
            self.assertEqual(3, len(records))
            for rec in records:
                self.assertEqual(base['labels']['仿真路径'], rec['labels']['仿真路径'])
            profile = json.loads((f.PACKAGE / 'check-profile.json').read_text(encoding='utf-8'))
            check = c.validate(f.PACKAGE / 'template.docx', out, profile)
            self.assertTrue(check['valid'], check)


if __name__ == '__main__':
    unittest.main()
